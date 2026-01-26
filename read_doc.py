# -*- coding: utf-8 -*-
from docx import Document

doc = Document(r'c:\Users\Laptop\Desktop\毕业设计\bishe\任务书.docx')

print("=== PARAGRAPHS ===")
for p in doc.paragraphs:
    if p.text.strip():
        print(p.text)

print("\n=== TABLES ===")
for i, table in enumerate(doc.tables):
    print(f"\n--- Table {i+1} ---")
    for row in table.rows:
        row_text = [cell.text.strip() for cell in row.cells]
        print(" | ".join(row_text))
